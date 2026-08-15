from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from .export_streaming import write_query_xlsx
from .raw_period_comparison_strict_ambiguity import AmbiguityAwareRawPeriodComparisonService
from .spreadsheet_utils import sanitize_excel_row


class ScalableRawPeriodComparisonService(AmbiguityAwareRawPeriodComparisonService):
    """Comparaison RAW stricte, compatible avec fusions et exports exhaustifs en streaming."""

    HEADERS = [
        "Statut","Matricule A","Matricule B","Nom A","Nom B","Prénom A","Prénom B",
        "Commun matricule","Commun nom","Régime A","Régime B","Institution A","Institution B",
        "Répétitions A","Répétitions B","Brut A","Brut B","Écart brut","Net A","Net B","Écart net",
        "Section A","Section B","Catégorie A","Catégorie B","Grade A","Grade B","Unité A",
        "Unité B","Province A","Province B","Diagnostic",
    ]

    def _result_query(self, status=""):
        condition = "comparaison_id=?"
        params = []
        special = {
            "MEME_MATRICULE_NOM_DIFFERENT": "meme_matricule_nom_different",
            "MEME_NOM_MATRICULE_DIFFERENT": "meme_nom_matricule_different",
            "DOUBLON_MATRICULE_A": "doublon_matricule_a",
            "DOUBLON_MATRICULE_B": "doublon_matricule_b",
            "DOUBLON_NOM_A": "doublon_nom_a",
            "DOUBLON_NOM_B": "doublon_nom_b",
        }
        if status in special:
            condition += f" AND {special[status]}"
        elif status:
            condition += " AND statut=?"
            params.append(status)
        query = f"""SELECT statut,matricule_a,matricule_b,nom_a,nom_b,prenom_a,prenom_b,
                commun_matricule,commun_nom,regime_a,regime_b,institution_a,institution_b,occurrences_a,occurrences_b,
                brut_a,brut_b,ecart_brut,net_a,net_b,ecart_net,section_a,section_b,categorie_a,categorie_b,
                grade_a,grade_b,unite_a,unite_b,province_a,province_b,diagnostic
            FROM resultats_comparaison_raw_periode
            WHERE {condition}
            ORDER BY commun_matricule DESC,commun_nom DESC,ABS(ecart_brut) DESC"""
        return query, params

    def _export_raw_complete(self, con, comparison_id: str, side: str, path: Path):
        src = con.execute(
            "SELECT table_source FROM sources_comparaison_raw_periode WHERE comparaison_id=? AND cote=? LIMIT 1",
            [comparison_id, side],
        ).fetchone()
        if not src:
            return 0
        table = src[0]
        safe = self._quote(table)
        columns = [r[0] for r in con.execute(f"DESCRIBE {safe}").fetchall()]
        ids = [r[0] for r in con.execute(
            "SELECT execution_id FROM sources_comparaison_raw_periode WHERE comparaison_id=? AND cote=? AND execution_id IS NOT NULL",
            [comparison_id, side],
        ).fetchall()]
        if "execution_id" in columns and ids:
            placeholders = ",".join("?" for _ in ids)
            query = f"SELECT * FROM {safe} WHERE execution_id IN ({placeholders})"
            params = ids
        else:
            query = f"SELECT * FROM {safe}"
            params = []
        return write_query_xlsx(con, path, query, params, columns, f"RAW_{side}")

    def export_all(self, comparison_id: str, parent_folder, progress=None):
        info = self.get_comparison(comparison_id)
        folder = Path(parent_folder) / f"comparaison_raw_{info['quarter']}_{info['year']}_{datetime.now():%Y%m%d_%H%M%S}"
        folder.mkdir(parents=True, exist_ok=True)
        progress and progress(5, "Création de la synthèse")

        base, metrics = self.summary(comparison_id)
        book = Workbook()
        sheet = book.active
        sheet.title = "Synthèse"
        sheet.append(["Indicateur", "Valeur"])
        for key, value in [
            ("Table A", info["table_a"]),
            ("Table B", info["table_b"]),
            ("Période", f"{info['quarter']} {info['year']}"),
            ("Communs par matricule", metrics[0] or 0),
            ("Communs par nom", metrics[1] or 0),
            ("Communs matricule + nom", metrics[2] or 0),
            ("Même matricule / nom différent", metrics[3] or 0),
            ("Même nom / matricule différent", metrics[4] or 0),
        ]:
            sheet.append([key, value])
        sheet.append([])
        sheet.append(["Statut", "Agents", "Brut A", "Brut B", "Net A", "Net B"])
        for row in base:
            sheet.append(list(sanitize_excel_row(row)))
        book.save(folder / "00_synthese.xlsx")

        exports = [
            ("01_tous_resultats.xlsx", ""),
            ("02_communs_matricule.xlsx", "COMMUN_PAR_MATRICULE"),
            ("03_communs_nom.xlsx", "COMMUN_PAR_NOM"),
            ("04_communs_matricule_et_nom.xlsx", "COMMUN_PAR_MATRICULE_ET_NOM"),
            ("05_uniquement_A.xlsx", "UNIQUEMENT_A"),
            ("06_uniquement_B.xlsx", "UNIQUEMENT_B"),
            ("07_meme_matricule_nom_different.xlsx", "MEME_MATRICULE_NOM_DIFFERENT"),
            ("08_meme_nom_matricule_different.xlsx", "MEME_NOM_MATRICULE_DIFFERENT"),
            ("09_doublons_matricule_A.xlsx", "DOUBLON_MATRICULE_A"),
            ("10_doublons_matricule_B.xlsx", "DOUBLON_MATRICULE_B"),
            ("11_doublons_nom_A.xlsx", "DOUBLON_NOM_A"),
            ("12_doublons_nom_B.xlsx", "DOUBLON_NOM_B"),
            ("13_matchs_ambigus_matricule.xlsx", "MATCH_AMBIGU_MATRICULE"),
            ("14_matchs_ambigus_nom.xlsx", "MATCH_AMBIGU_NOM"),
        ]
        with self.db.connect() as con:
            for index, (filename, status) in enumerate(exports, start=1):
                progress and progress(8 + int(62 * index / len(exports)), f"Export complet {filename}")
                query, extra = self._result_query(status)
                write_query_xlsx(con, folder / filename, query, [comparison_id] + extra, self.HEADERS)

            progress and progress(76, "Annexe RAW A complète")
            self._export_raw_complete(con, comparison_id, "A", folder / "15_annexe_RAW_A_complete.xlsx")
            progress and progress(88, "Annexe RAW B complète")
            self._export_raw_complete(con, comparison_id, "B", folder / "16_annexe_RAW_B_complete.xlsx")
            con.execute(
                "UPDATE comparaisons_raw_periode SET dossier_export=? WHERE comparaison_id=?",
                [str(folder), comparison_id],
            )

        progress and progress(100, "Export exhaustif terminé")
        return str(folder)
