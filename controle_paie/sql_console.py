from __future__ import annotations

import os
import re
import time
from pathlib import Path

from openpyxl import Workbook

from .export_streaming import append_query_sheets, atomic_save_workbook, write_query_csv


class SqlConsoleService:
    """Console SQL DuckDB en lecture seule pour exploration et export."""

    ALLOWED_PREFIXES = {"SELECT", "WITH", "DESCRIBE", "DESC", "EXPLAIN", "SHOW"}
    BLOCKED_KEYWORDS = {
        "INSERT", "UPDATE", "DELETE", "DROP", "ALTER", "TRUNCATE", "CREATE", "REPLACE",
        "MERGE", "COPY", "ATTACH", "DETACH", "INSTALL", "LOAD", "CALL", "EXPORT", "IMPORT",
        "VACUUM", "CHECKPOINT", "SET", "RESET", "GRANT", "REVOKE", "PRAGMA",
    }

    def __init__(self, db):
        self.db = db

    def list_raw_tables(self) -> list[tuple[str, int]]:
        with self.db.connect() as con:
            rows = con.execute("""SELECT table_name FROM information_schema.tables
                WHERE table_schema='main' AND table_name LIKE 'raw_%' ORDER BY table_name""").fetchall()
            result = []
            for (name,) in rows:
                safe = str(name).replace('"', '""')
                count = int(con.execute(f'SELECT COUNT(*) FROM "{safe}"').fetchone()[0])
                result.append((str(name), count))
        return result

    def describe_table(self, table_name: str) -> list[tuple]:
        if table_name not in {name for name, _ in self.list_raw_tables()}:
            raise ValueError("Table RAW introuvable.")
        safe = table_name.replace('"', '""')
        with self.db.connect() as con:
            return con.execute(f'DESCRIBE "{safe}"').fetchall()

    def sample_table(self, table_name: str, limit: int = 20) -> tuple[list[str], list[tuple]]:
        if table_name not in {name for name, _ in self.list_raw_tables()}:
            raise ValueError("Table RAW introuvable.")
        safe = table_name.replace('"', '""')
        limit = max(1, min(int(limit), 200))
        with self.db.connect() as con:
            cursor = con.execute(f'SELECT * FROM "{safe}" LIMIT ?', [limit])
            columns = [item[0] for item in cursor.description]
            rows = cursor.fetchall()
        return columns, rows

    @classmethod
    def validate_read_only_query(cls, query: str) -> str:
        text = (query or "").strip()
        if not text:
            raise ValueError("Saisissez une requête SQL.")
        cleaned = re.sub(r"--[^\n]*", " ", text)
        cleaned = re.sub(r"/\*.*?\*/", " ", cleaned, flags=re.S).strip()
        if not cleaned:
            raise ValueError("La requête ne contient aucune instruction SQL.")
        statements = [part.strip() for part in cleaned.split(";") if part.strip()]
        if len(statements) != 1:
            raise ValueError("Une seule instruction SQL peut être exécutée à la fois.")
        statement = statements[0]
        first = re.match(r"([A-Za-z_]+)", statement)
        keyword = first.group(1).upper() if first else ""
        if keyword not in cls.ALLOWED_PREFIXES:
            raise ValueError("Seules les requêtes de lecture SELECT, WITH, DESCRIBE, EXPLAIN et SHOW sont autorisées.")
        tokens = {token.upper() for token in re.findall(r"\b[A-Za-z_]+\b", statement)}
        blocked = sorted(tokens.intersection(cls.BLOCKED_KEYWORDS))
        if blocked:
            raise ValueError("Instruction interdite en mode lecture seule : " + ", ".join(blocked) + ".")
        return statement

    def execute(self, query: str, display_limit: int = 1000) -> dict:
        statement = self.validate_read_only_query(query)
        display_limit = max(1, min(int(display_limit), 10000))
        started = time.perf_counter()
        with self.db.connect() as con:
            cursor = con.execute(statement)
            columns = [item[0] for item in (cursor.description or [])]
            rows = cursor.fetchmany(display_limit + 1) if columns else []
        elapsed = time.perf_counter() - started
        truncated = len(rows) > display_limit
        if truncated:
            rows = rows[:display_limit]
        return {
            "query": statement,
            "columns": columns,
            "rows": rows,
            "displayed": len(rows),
            "truncated": truncated,
            "elapsed": elapsed,
        }

    def export_csv(self, query: str, path: str | Path) -> Path:
        statement = self.validate_read_only_query(query)
        target = Path(path)
        with self.db.connect() as con:
            write_query_csv(con, target, statement)
        return target

    def export_excel(self, query: str, path: str | Path) -> Path:
        statement = self.validate_read_only_query(query)
        target = Path(path)
        book = Workbook(write_only=True)
        with self.db.connect() as con:
            count = append_query_sheets(book, con, statement, sheet_name="Résultats")
        meta = book.create_sheet("Requête SQL")
        meta.append(["Requête"])
        meta.append([statement])
        meta.append(["Nombre de lignes", count])
        atomic_save_workbook(book, target)
        return target

    def export_parquet(self, query: str, path: str | Path) -> Path:
        """Export natif DuckDB, adapté aux résultats très volumineux."""
        statement = self.validate_read_only_query(query)
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_name(f".{target.name}.part")
        escaped = str(temporary).replace("'", "''")
        try:
            with self.db.connect() as con:
                con.execute(f"COPY ({statement}) TO '{escaped}' (FORMAT PARQUET, COMPRESSION ZSTD)")
            os.replace(temporary, target)
        finally:
            if temporary.exists():
                temporary.unlink()
        return target
